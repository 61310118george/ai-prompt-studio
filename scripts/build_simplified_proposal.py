"""Build a concise, visual-first version of the capstone proposal."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "word" / "AI提示詞視覺化管理App_完整企劃書_精簡易讀版.docx"
EVIDENCE = ROOT / "docs" / "evidence" / "revision-21"

FONT_LATIN = "Arial"
FONT_CJK = "PingFang TC"
INK = "1F1F1F"
MUTED = "666666"
ACCENT = "8B513A"
ACCENT_LIGHT = "F2E7DF"
BLUE = "355C7D"
BLUE_LIGHT = "EAF1F7"
PALE = "F7F7F7"
BORDER = "D9D9D9"


def set_run(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = FONT_LATIN
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, top=130, start=150, bottom=130, end=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table, color=BORDER, size="5"):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.first_child_found_in("w:tblBorders")
    if existing is not None:
        tbl_pr.remove(existing)
    container = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)
        container.append(node)
    tbl_pr.append(container)


def clear_borders(table):
    borders(table, color="FFFFFF", size="0")


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run(run, 9, color=MUTED)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    settings = {
        "Title": (25, True, 0, 10),
        "Subtitle": (13.5, False, 0, 14),
        "Heading 1": (18, True, 9, 7),
        "Heading 2": (13.5, True, 8, 4),
        "Heading 3": (11.5, True, 6, 3),
    }
    for name, (size, bold, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    header = section.header.paragraphs[0]
    header.text = "AI 提示詞視覺化管理 App 專題企劃書"
    for run in header.runs:
        set_run(run, 8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(footer.add_run("第 "), 9, color=MUTED)
    add_page_number(footer)
    set_run(footer.add_run(" 頁"), 9, color=MUTED)

    doc.core_properties.title = "AI 提示詞視覺化管理 App 專題企劃書 精簡易讀版"
    doc.core_properties.subject = "AI Prompt Studio 專題企劃書"
    doc.core_properties.author = "AI Prompt Studio 專題小組"


def add_text(doc, text, *, size=11, bold=False, color=INK, align=None, before=0, after=5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), size, bold, color)
    return p


def add_bullets(doc, items, *, size=10.8, compact=False):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(2 if compact else 4)
        set_run(p.add_run(item), size)


def add_picture(doc, filename, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(EVIDENCE / filename), width=Inches(width))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    set_run(p.add_run(caption), 9.3, color=MUTED)


def add_table(doc, headers, rows, widths=None, font_size=9.5, header_fill=BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False if widths else True
    borders(table)
    for idx, label in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Cm(widths[idx])
        shade(cell, header_fill)
        margins(cell, 150, 140, 150, 140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(label), font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            if widths:
                cell.width = Cm(widths[idx])
            shade(cell, PALE if row_index % 2 else "FFFFFF")
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(value) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_run(p.add_run(value), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_flow(doc, steps):
    table = doc.add_table(rows=1, cols=len(steps) * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    clear_borders(table)
    for index, step in enumerate(steps):
        cell = table.cell(0, index * 2)
        cell.width = Cm(2.75)
        shade(cell, ACCENT_LIGHT if index % 2 == 0 else BLUE_LIGHT)
        margins(cell, 180, 120, 180, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(f"{index + 1}\n"), 13, True, ACCENT)
        set_run(p.add_run(step), 9.5, True)
        if index < len(steps) - 1:
            arrow = table.cell(0, index * 2 + 1)
            arrow.width = Cm(0.6)
            arrow.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = arrow.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run(p.add_run("→"), 14, True, MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    configure(doc)

    # Cover
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("AI 提示詞視覺化管理 App"), 25, True)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("專題企劃書 精簡易讀版"), 14, False, ACCENT)
    add_text(doc, "把散落的 Prompt 集中管理，並在本機查看版本、Token、專案記憶與選用 AI 分析。",
             size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(EVIDENCE / "prompts-1440.png"), width=Inches(5.7))
    add_text(doc, "學生：____________　系級：____________　指導老師：____________",
             size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=3)
    add_text(doc, "版本 V1.4　2026 年 10 月", size=9.5, color=MUTED,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    page_break(doc)

    # Quick overview
    doc.add_heading("一分鐘看懂這個專題", level=1)
    add_text(doc, "這是一套給學生與個人開發者使用的本機工具。它把 Prompt、Markdown 與 AI 記憶檔集中在專案中，減少找不到版本、檔案作用不清楚與修改後無法復原的問題。", size=12, after=9)
    add_table(doc, ["使用者現在遇到的問題", "App 提供的做法"], [
        ("Prompt 散落在聊天與筆記裡", "依專案分類、搜尋並保存版本"),
        ("不知道 Prompt 是否缺少必要資訊", "用範本、草稿與版本引導整理"),
        ("內容變長，卻看不到輸入負擔", "在本機顯示 Token 與用量等級"),
        ("AGENTS.md 等記憶檔不容易理解", "顯示檔案用途、資料夾分組與健檢問題"),
    ], widths=[7.8, 9.2], font_size=10)
    doc.add_heading("核心操作", level=2)
    add_flow(doc, ["選擇專案", "整理 Prompt", "檢查負擔", "保存版本"])
    add_text(doc, "核心功能離線運作，不需要 API Key。Gemini 分析是選用功能，只有使用者確認後才會送出目前草稿。", size=10.5, color=MUTED, before=5)
    page_break(doc)

    # Audience and goals
    doc.add_heading("使用者與使用情境", level=1)
    add_table(doc, ["使用者", "最常遇到的情況", "App 能幫上什麼"], [
        ("學生", "研究、報告與課業 Prompt 很多", "依課程或題目分類，快速找回最後版本"),
        ("個人開發者", "程式審查與開發指令重複使用", "比較版本、Token 與 Markdown 輸出"),
        ("小型團隊成員", "專案規則與 Agent 記憶不透明", "查看檔案作用域、用途與健檢問題"),
    ], widths=[3.2, 6.8, 7.0], font_size=9.8)
    doc.add_heading("V1.4 要做到的五件事", level=2)
    add_bullets(doc, [
        "把 Prompt 依專案集中管理，支援搜尋、標籤與版本。",
        "用內建範本與自己的範本協助建立草稿。",
        "顯示純文字 Token 與輸入負擔，不宣稱等於實際帳單。",
        "依核心記憶、專案提示詞、專案技能與 AI 輸出項目分組。",
        "安全管理 Markdown，保留差異預覽、備份與外部修改檢查。",
    ])
    doc.add_heading("不會做的事", level=2)
    add_text(doc, "V1.4 不提供多人雲端協作，不保證 AI 回答品質提升，也不會自動套用 Gemini 分析結果。所有檔案寫入仍需使用者預覽與確認。", size=10.8)
    page_break(doc)

    # Prompt workbench
    doc.add_heading("主要功能一 Prompt 工作台", level=1)
    add_picture(doc, "prompts-1440.png", "圖 1　提示詞工作台整合空白草稿、內建範本、已儲存提示詞與輸入負擔提示。", width=6.45)
    add_table(doc, ["功能", "用途"], [
        ("範本與群組", "建立空白草稿、套用內建範本，自己的範本可拖曳分類"),
        ("版本紀錄", "每次儲存建立新版本，可帶回舊內容再編輯"),
        ("輸入負擔", "顯示 Token 與低、中、高等級，不把估算當完整帳單"),
        ("選用 AI", "使用者同意後才送出草稿給 Gemini 分析"),
    ], widths=[4.2, 12.8], font_size=10)
    page_break(doc)

    # Project map and memory programming
    doc.add_heading("主要功能二 專案檔案與章節編輯", level=1)
    add_picture(doc, "projects-1440.png", "圖 2　專案檔案依核心記憶、專案提示詞、專案技能與 AI 輸出項目分組。", width=5.9)
    add_picture(doc, "editor-1440.png", "圖 3　編輯器可改章節、新增章節或進階全文修改，儲存前先看差異。", width=5.9)
    add_text(doc, "一般 Markdown 可以查看與編輯，但不會被誤標成 Agent 自動記憶。企劃書、規格、展示指南與交付紀錄會集中到 AI 輸出項目。", size=10.5, color=MUTED)
    page_break(doc)

    # Health and sources
    doc.add_heading("主要功能三 記憶健檢與規格來源", level=1)
    add_picture(doc, "health-1440.png", "圖 4　健檢依資料夾分組，右側以卡片顯示行號、類型與問題內容。", width=5.9)
    add_picture(doc, "settings-1440.png", "圖 5　設定頁顯示版本、API 狀態、Gemini Key 設定與規格來源。", width=5.9)
    add_text(doc, "健檢只提出問題，不會自動修改檔案。Gemini API Key 只保存在目前 macOS 使用者的 App Support 私有檔。", size=10.5, color=MUTED)
    page_break(doc)

    # Architecture and safety
    doc.add_heading("系統怎麼運作", level=1)
    add_table(doc, ["層級", "使用技術", "負責工作"], [
        ("介面", "HTML、CSS、JavaScript", "Web 與 macOS App 共用同一套畫面"),
        ("桌面橋接", "pywebview", "讓 Web UI 呼叫本機 Python 功能"),
        ("服務", "Python", "掃描專案、Token 計數、健檢、Diff 與檔案寫回"),
        ("資料", "SQLite、Markdown、本機私有設定", "保存 Prompt、版本、群組、專案檔案與 API Key 狀態"),
    ], widths=[3.0, 5.0, 9.0], font_size=9.8)
    doc.add_heading("安全寫回流程", level=2)
    add_flow(doc, ["產生預覽", "檢查 Diff", "比對雜湊", "備份後寫入"])
    add_bullets(doc, [
        "檔案若被其他程式修改，App 會停止套用舊預覽。",
        "寫入採暫存檔與 atomic replace，避免只寫到一半。",
        "只能修改使用者選定根目錄內的 Markdown。",
        "每次寫入前建立備份，之後可以查看紀錄或復原。",
    ], compact=True)
    doc.add_heading("Token 與成本怎麼看", level=2)
    add_text(doc, "App 顯示選定 tokenizer 的純文字 Token 與輸入負擔等級。它沒有連接帳號額度、重置時間或完整帳單，因此不能宣稱一定省錢或一定夠用。", size=10.5)
    page_break(doc)

    # Evaluation
    doc.add_heading("如何證明這個 App 有用", level=1)
    add_text(doc, "開發完成不代表研究問題已經成立。專題會分別驗證功能、Token 變化與實際使用效果。", size=11.5, after=8)
    add_table(doc, ["要驗證什麼", "怎麼測", "判斷方式"], [
        ("功能是否可靠", "測試資料庫、版本、Diff、備份、復原與路徑安全", "自動測試全部通過"),
        ("Token 是否減少", "比較原文與候選，記錄沒有改善的樣本", "平均值、中位數、範圍與零改善比例"),
        ("內容品質是否保留", "固定模型與條件，比較原文和候選輸出", "任務成功、事實保留、格式與限制"),
        ("管理是否更容易", "6 至 10 位受試者完成找 Prompt 任務", "時間、成功率、找錯版本與主觀難度"),
    ], widths=[4.0, 7.2, 5.8], font_size=9.5)
    doc.add_heading("研究順序", level=2)
    add_flow(doc, ["功能測試", "Token 測量", "真人試用", "品質比較"])
    add_text(doc, "目前已完成工程測試與 V1.4 Web/桌面 smoke test。真人可用性研究及模型輸出品質比較尚未執行，因此文件不宣稱已證明節省比例或品質等效。", size=10.5, color=MUTED, before=5)
    page_break(doc)

    # Scope, status and plan
    doc.add_heading("目前進度與後續安排", level=1)
    add_table(doc, ["狀態", "內容"], [
        ("已完成", "共用 Web UI、macOS App、Prompt 工作台、群組、章節編輯、健檢、Gemini 設定與安全寫回"),
        ("已驗證", "52 項 Python 測試、V1.4 Web/桌面 smoke test、操作截圖、Word 與 PDF 文件"),
        ("待完成", "真人可用性研究、真實 Gemini API 帳號驗證、模型品質比較、正式報告與展示影片"),
    ], widths=[3.5, 13.5], font_size=10)
    doc.add_heading("八週建議時程", level=2)
    add_table(doc, ["週次", "工作", "交付內容"], [
        ("1", "確認題目與研究方法", "企劃與 Spec 定稿"),
        ("2 至 3", "整理功能與回歸測試", "V1.4 可操作 App"),
        ("4", "Token 基準與錯誤案例", "可重現數據"),
        ("5 至 6", "品質比較與可用性試測", "原始數據與限制"),
        ("7", "修正問題與完整驗收", "測試版本"),
        ("8", "報告、簡報與展示", "期末交付包"),
    ], widths=[2.5, 7.0, 7.5], font_size=9.5)
    doc.add_heading("主要風險", level=2)
    add_bullets(doc, [
        "精簡可能刪掉必要語意，因此必須保留 Diff 與人工確認。",
        "Token 不等於完整帳單，因此畫面要標明計算範圍。",
        "真實樣本不足時，不能把合成測試結果當成普遍結論。",
        "功能範圍若擴大為企業平台，三個月內不容易完成。",
    ], compact=True)
    page_break(doc)

    # Deliverables and teacher decisions
    doc.add_heading("預期交付與老師需要確認的事項", level=1)
    doc.add_heading("預期交付", level=2)
    add_table(doc, ["類型", "內容"], [
        ("軟體", "macOS App、Web UI、原始碼與測試"),
        ("資料", "Token 基準、研究紀錄與驗收證據"),
        ("文件", "企劃書、完整 Spec、操作說明與 Demo 流程"),
    ], widths=[3.5, 13.5], font_size=10)
    doc.add_heading("請老師確認", level=2)
    add_bullets(doc, [
        "正式題目名稱與學校文件格式。",
        "工程成果與研究成果的評分比重。",
        "可用性試測是否需要校內研究倫理程序。",
        "模型品質比較的評分標準與可用預算。",
        "正式期限、組員分工與展示方式。",
    ])
    doc.add_heading("給老師的簡短說明", level=2)
    add_text(doc, "本專題以專案為單位管理 Prompt，提供範本、搜尋、版本、Token 輸入負擔、Markdown 章節編輯與記憶健檢，也能選用 Gemini 做草稿分析。目前已有 V1.4 可操作版本，下一步會用實驗確認整理效率、Token 變化與輸出品質。", size=11)

    # References
    doc.add_heading("參考資料", level=1)
    add_text(doc, "以下資料用於企劃結構、需求規格、Prompt 管理、Token 計數與 API 安全邊界。查閱日期為 2026 年 10 月 8 日。", size=10.5, color=MUTED)
    refs = [
        "S1 Langfuse Prompt Management",
        "S2 Promptfoo Introduction",
        "S3 OpenAI tiktoken",
        "S4 Gemini API Pricing",
        "S5 Gemini API Key Guide",
        "S6 OWASP Prompt Injection Prevention Cheat Sheet",
        "S7 ISO IEC IEEE 29148 2018 公開介紹",
    ]
    add_bullets(doc, refs, size=10.2, compact=True)
    doc.add_heading("一句話總結", level=2)
    add_text(doc, "AI Prompt Studio Next V1.4 把 Prompt、Markdown 與 Agent 記憶整理成可以搜尋、分類、預覽、健檢、備份與復原的本機工作流程。", size=12, bold=True, color=ACCENT)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
