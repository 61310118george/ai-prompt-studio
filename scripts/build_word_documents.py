"""Build editable Word versions of the illustrated proposal and SRS."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "word"
OUTPUT.mkdir(parents=True, exist_ok=True)
FONT_LATIN = "Arial"
FONT_CJK = "PingFang TC"
INK = "000000"
ACCENT = "8B513A"
HEADER_FILL = "6C493B"
ALT_FILL = "F6F1EC"
BORDER = "D9D9D9"


def set_run_font(run, size=None, bold=None, color=INK):
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), BORDER)
        borders.append(node)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run_font(run, 9, color="666666")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text):
    pattern = re.compile(r"\[([^]]+)]\((https?://[^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, 11)
        if match.group(1):
            add_hyperlink(paragraph, match.group(1), match.group(2))
        elif match.group(3):
            run = paragraph.add_run(match.group(3))
            set_run_font(run, 10)
            run.font.name = "Courier New"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Courier New")
        else:
            run = paragraph.add_run(match.group(4))
            set_run_font(run, 11, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, 11)


def clean_heading(text):
    text = text.replace("：", " ").replace(":", " ").replace("、", " ")
    text = re.sub(r"(?<=\d)\.(?=\d)", " ", text)
    text = text.replace("-", " ")
    return re.sub(r"\s+", " ", text).strip()


def clear_paragraph_borders(paragraph_or_style):
    """Remove template-provided decorative rules from title paragraphs/styles."""
    element = paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def configure_document(doc, title):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, space_before, space_after in (("Title", 25, 0, 12), ("Subtitle", 14, 0, 20),
                                                        ("Heading 1", 18, 14, 8), ("Heading 2", 14, 11, 6),
                                                        ("Heading 3", 12, 8, 4)):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            clear_paragraph_borders(style)

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = title
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        set_run_font(run, 8.5, color="666666")

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, 9, color="666666")
    add_page_field(paragraph)
    suffix = paragraph.add_run(" 頁")
    set_run_font(suffix, 9, color="666666")

    props = doc.core_properties
    props.title = title
    props.subject = "可編輯的專題文件"
    props.author = "AI Prompt Studio 專題小組"
    props.keywords = "提示詞管理, Token, Agent 記憶, 專題"


def add_cover(doc, title, subtitle, version_line):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    clear_paragraph_borders(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, 25, bold=True)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, 14, color=ACCENT)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(version_line)
    set_run_font(run, 10.5, color="555555")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("學生：____________　系級：____________")
    set_run_font(run, 11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("指導老師：____________")
    set_run_font(run, 11)
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(6)
    table.columns[1].width = Cm(6)
    for index, value in enumerate(("專題文件", "2026 10 08")):
        cell = table.cell(0, index)
        set_cell_shading(cell, "F0E4D9")
        set_cell_margins(cell, 180, 160, 180, 160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(value)
        set_run_font(run, 10, color="555555")
    set_table_borders(table)
    doc.add_page_break()


def add_contents(doc, headings):
    doc.add_heading("目錄", level=1)
    p = doc.add_paragraph()
    run = p.add_run("下列章節可在 Word 導覽窗格中直接跳轉。修改標題後，這份清單可一併手動調整。")
    set_run_font(run, 10, color="666666")
    for heading in headings:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.7)
        run = p.add_run(clean_heading(heading))
        set_run_font(run, 11)
    doc.add_page_break()


def add_markdown_table(doc, lines):
    parsed = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        parsed.append(cells)
    if not parsed:
        return
    cols = max(len(row) for row in parsed)
    table = doc.add_table(rows=len(parsed), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, values in enumerate(parsed):
        row = table.rows[row_index]
        if row_index == 0:
            set_repeat_table_header(row)
        for col_index in range(cols):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, HEADER_FILL if row_index == 0 else (ALT_FILL if row_index % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            text = values[col_index] if col_index < len(values) else ""
            add_inline(p, text)
            for run in p.runs:
                set_run_font(run, 9 if cols >= 4 else 9.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_body(doc, source):
    lines = source.read_text(encoding="utf-8").splitlines()
    start = next((index for index, line in enumerate(lines) if line.startswith("## ")), 0)
    lines = lines[start:]
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("```"):
            index += 1
            code = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.6)
            run = p.add_run("\n".join(code))
            set_run_font(run, 9)
            run.font.name = "Courier New"
            continue
        image_match = re.fullmatch(r"!\[([^]]*)]\(([^)]+)\)", stripped)
        if image_match:
            path = (source.parent / image_match.group(2)).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Inches(6.35))
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        heading = re.match(r"^(#{2,3})\s+(.+)$", stripped)
        if heading:
            level = 1 if len(heading.group(1)) == 2 else 2
            p = doc.add_heading(clean_heading(heading.group(2)), level=level)
            index += 1
            continue
        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            add_inline(p, re.sub(r"^[-*]\s+", "", stripped))
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.7)
            add_inline(p, numbered.group(2))
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line or re.match(r"^#{2,3}\s+", next_line) or next_line.startswith("|") or next_line.startswith("```") or re.match(r"^[-*]\s+", next_line) or re.match(r"^\d+\.\s+", next_line) or re.fullmatch(r"!\[[^]]*]\([^)]+\)", next_line):
                break
            paragraph_lines.append(next_line)
            index += 1
        text = " ".join(paragraph_lines)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if re.match(r"^圖\s*(?:S-)?\d+", text):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(9)
            add_inline(p, text)
            for run in p.runs:
                set_run_font(run, 9.5, color="555555")
        else:
            add_inline(p, text)


def build(source_name, output_name, title, subtitle, version):
    source = ROOT / "docs" / source_name
    headings = [match.group(1) for line in source.read_text(encoding="utf-8").splitlines()
                if (match := re.match(r"^##\s+(.+)$", line))]
    doc = Document()
    configure_document(doc, title)
    add_cover(doc, title, subtitle, version)
    add_contents(doc, headings)
    add_body(doc, source)
    target = OUTPUT / output_name
    doc.save(target)
    print(target)


if __name__ == "__main__":
    build("專題企劃書.md", "AI提示詞視覺化管理App_完整企劃書.docx",
          "AI 提示詞視覺化管理 App",
          "本機提示詞管理 專案記憶與 Token 輸入負擔",
          "版本 V1.4　精簡企劃書與操作介面")
    build("SPEC.md", "AI提示詞視覺化管理App_軟體需求規格書.docx",
          "AI Prompt Studio",
          "軟體需求與技術規格書",
          "版本 V1.4　功能需求 介面 資料 API 安全與驗收")
